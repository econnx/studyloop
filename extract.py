"""Text extraction from PDF, PPTX, DOCX, HTML, text files and images."""
import io
from html.parser import HTMLParser

import docx
import pymupdf
from pptx import Presentation

IMAGES = (".png", ".jpg", ".jpeg", ".webp")
SUPPORTED = (".pdf", ".pptx", ".docx", ".txt", ".md", ".html", ".htm") + IMAGES
MIN_PAGE_TEXT = 50  # a PDF page with fewer characters is treated as a scan


def is_supported(filename: str) -> bool:
    return filename.lower().endswith(SUPPORTED)


class _HTMLText(HTMLParser):
    SKIP = {"style", "script", "noscript", "svg"}
    BLOCK = {"p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6", "section", "article", "table"}

    def __init__(self):
        super().__init__()
        self.parts, self.skip = [], 0

    def handle_starttag(self, tag, attrs):
        if tag in self.SKIP:
            self.skip += 1
        elif tag in self.BLOCK:
            self.parts.append("\n")
        elif tag == "img" and (alt := dict(attrs).get("alt")):
            self.parts.append(f" [Image: {alt}] ")

    def handle_endtag(self, tag):
        if tag in self.SKIP:
            self.skip = max(0, self.skip - 1)
        elif tag in self.BLOCK:
            self.parts.append("\n")

    def handle_data(self, data):
        if not self.skip:
            self.parts.append(data)


def html_to_text(html: str) -> str:
    p = _HTMLText()
    p.feed(html)
    lines = (" ".join(line.split()) for line in "".join(p.parts).splitlines())
    return "\n".join(line for line in lines if line)


def extract_text(filename: str, data: bytes, read_image=None, pdf_as_images: bool = False) -> str:
    """read_image(png_bytes) -> str recognizes images with the AI (if None, images are skipped)."""
    name = filename.lower()
    if name.endswith(IMAGES):
        return read_image(data) if read_image else ""
    if name.endswith(".pdf"):
        pages = []
        with pymupdf.open(stream=data, filetype="pdf") as doc:
            for i, page in enumerate(doc, 1):
                text = page.get_text()
                if read_image and (pdf_as_images or len(text.strip()) < MIN_PAGE_TEXT):
                    text = read_image(page.get_pixmap(dpi=150).tobytes("png"))
                pages.append(f"--- Page {i} ---\n{text}")
        return "\n\n".join(pages)
    if name.endswith(".pptx"):
        prs = Presentation(io.BytesIO(data))
        slides = []
        for i, slide in enumerate(prs.slides, 1):
            parts = [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
                parts.append("Notes: " + slide.notes_slide.notes_text_frame.text)
            if parts:
                slides.append(f"--- Slide {i} ---\n" + "\n".join(parts))
        return "\n\n".join(slides)
    if name.endswith(".docx"):
        d = docx.Document(io.BytesIO(data))
        return "\n".join(p.text for p in d.paragraphs if p.text.strip())
    if name.endswith((".html", ".htm")):
        return html_to_text(data.decode("utf-8", errors="ignore"))
    if name.endswith((".txt", ".md")):
        return data.decode("utf-8", errors="ignore")
    raise ValueError(f"Unsupported format: {filename}")
